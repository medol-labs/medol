from pathlib import Path
import re
import subprocess
import tempfile

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parents[1]
MD_DIR = ROOT / "tmp" / "联邦学习子系统文档"
DOCX_DIR = ROOT / "tmp" / "联邦学习子系统Word文档"
DOCX_DIR.mkdir(parents=True, exist_ok=True)

BODY_FONT = "STSong"
HEADING_FONT = "Hiragino Sans GB"


def main():
    for markdown_path in sorted(MD_DIR.glob("3.*-联邦学习子系统.md")):
        output_path = DOCX_DIR / f"{markdown_path.stem}.docx"
        build_docx(markdown_path, output_path)
        print(output_path)


def build_docx(markdown_path: Path, output_path: Path):
    doc = Document()
    configure_document(doc)
    parse_markdown_into_doc(doc, markdown_path.read_text(encoding="utf-8"))
    doc.save(output_path)


def configure_document(doc: Document):
    section = doc.sections[0]
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(2.6)

    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(10.5)
    set_style_font(normal, BODY_FONT)
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.space_after = Pt(4)

    for level in range(1, 7):
        style = doc.styles[f"Heading {level}"]
        style.font.name = HEADING_FONT
        set_style_font(style, HEADING_FONT)
        style.font.bold = level <= 3
        style.paragraph_format.space_before = Pt(10 if level <= 2 else 6)
        style.paragraph_format.space_after = Pt(5)
        if level == 1:
            style.font.size = Pt(18)
            style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif level == 2:
            style.font.size = Pt(15)
        elif level == 3:
            style.font.size = Pt(13)
        else:
            style.font.size = Pt(11)


def parse_markdown_into_doc(doc: Document, markdown: str):
    lines = markdown.splitlines()
    index = 0
    in_code = False
    code_language = ""
    code_lines = []

    while index < len(lines):
        line = lines[index].rstrip()

        if line.startswith("```"):
            if in_code:
                if code_language == "mermaid":
                    add_mermaid_diagram(doc, code_lines)
                else:
                    add_code_block(doc, code_lines)
                code_lines = []
                in_code = False
                code_language = ""
            else:
                in_code = True
                code_language = line.removeprefix("```").strip().lower()
            index += 1
            continue

        if in_code:
            code_lines.append(line)
            index += 1
            continue

        if not line.strip():
            index += 1
            continue

        heading = re.match(r"^(#{1,6})\s+(.+)$", line)
        if heading:
            level = min(6, len(heading.group(1)))
            text = clean_inline(heading.group(2))
            if level == 1:
                paragraph = doc.add_heading(text, level=1)
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                doc.add_heading(text, level=level)
            index += 1
            continue

        if is_table_start(lines, index):
            rows = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                if not is_table_separator(lines[index]):
                    rows.append(parse_table_row(lines[index]))
                index += 1
            add_table(doc, rows)
            continue

        if line.lstrip().startswith("- "):
            paragraph = doc.add_paragraph(style="List Bullet")
            add_run(paragraph, clean_inline(line.lstrip()[2:]))
            index += 1
            continue

        paragraph = doc.add_paragraph()
        add_run(paragraph, clean_inline(line))
        index += 1


def is_table_start(lines, index):
    return (
        index + 1 < len(lines)
        and lines[index].strip().startswith("|")
        and is_table_separator(lines[index + 1])
    )


def is_table_separator(line):
    stripped = line.strip()
    return bool(re.match(r"^\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)+\|?$", stripped))


def parse_table_row(line):
    cells = line.strip().strip("|").split("|")
    return [clean_inline(cell.strip()) for cell in cells]


def add_table(doc: Document, rows):
    if not rows:
        return
    column_count = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=column_count)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = True

    for row_index, table_row in enumerate(table.rows):
        row = rows[row_index]
        for column_index, cell in enumerate(table_row.cells):
            text = row[column_index] if column_index < len(row) else ""
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_shading(cell, "D9EAF7" if row_index == 0 else "FFFFFF")
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if row_index == 0 else WD_ALIGN_PARAGRAPH.LEFT
            run = paragraph.add_run(text)
            run.font.name = BODY_FONT
            run.font.size = Pt(8.5 if column_count >= 5 else 9)
            run.font.bold = row_index == 0
            set_run_font(run, BODY_FONT)

    doc.add_paragraph()


def add_code_block(doc: Document, lines):
    if not lines:
        return
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Cm(0.5)
    run = paragraph.add_run("\n".join(lines))
    run.font.name = "Courier New"
    run.font.size = Pt(8)


def add_mermaid_diagram(doc: Document, lines):
    nodes, edges = parse_mermaid_flowchart(lines)
    if not nodes and not edges:
        add_code_block(doc, lines)
        return

    image_path = render_topology_image(nodes, edges)
    if image_path:
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run()
        run.add_picture(str(image_path), width=Cm(15.0))

        caption = doc.add_paragraph()
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_run(caption, "图：联邦学习子系统部署拓扑")

    relation_rows = [["来源节点", "目标节点", "连接关系"]]
    for source, target in edges:
        relation_rows.append([nodes.get(source, source), nodes.get(target, target), "网络调用或数据交互"])
    add_table(doc, relation_rows)


def parse_mermaid_flowchart(lines):
    nodes = {}
    edges = []
    node_pattern = re.compile(r'^\s*([^\s\[]+)\["(.+)"\]\s*$')
    edge_pattern = re.compile(r"^\s*([^\s-]+)\s*-->\s*([^\s]+)\s*$")

    for line in lines:
        node_match = node_pattern.match(line)
        if node_match:
            nodes[node_match.group(1)] = node_match.group(2)
            continue

        edge_match = edge_pattern.match(line)
        if edge_match:
            source, target = edge_match.group(1), edge_match.group(2)
            edges.append((source, target))
            nodes.setdefault(source, source)
            nodes.setdefault(target, target)

    return nodes, edges


def render_topology_image(nodes, edges):
    if not nodes:
        return None

    with tempfile.TemporaryDirectory() as directory:
        dot_path = Path(directory) / "topology.dot"
        image_path = Path(directory) / "topology.png"
        node_ids = {key: f"n{index}" for index, key in enumerate(nodes.keys())}
        dot_lines = [
            "digraph G {",
            "  rankdir=TB;",
            '  graph [bgcolor="white", margin="0.18", nodesep="0.45", ranksep="0.55"];',
            '  node [shape=box, style="rounded,filled", color="#4F81BD", fillcolor="#EAF3F8", fontname="Arial Unicode MS", fontsize=14, margin="0.14,0.08"];',
            '  edge [color="#4F81BD", penwidth=1.4, arrowsize=0.8];',
        ]
        for key, label in nodes.items():
            dot_lines.append(f'  {node_ids[key]} [label="{escape_dot_label(label)}"];')
        for source, target in edges:
            dot_lines.append(f"  {node_ids[source]} -> {node_ids[target]};")
        dot_lines.append("}")
        dot_path.write_text("\n".join(dot_lines), encoding="utf-8")

        try:
            subprocess.run(
                ["dot", "-Tpng", str(dot_path), "-o", str(image_path)],
                check=True,
                stdout=subprocess.PIPE,
                stderr=subprocess.PIPE,
            )
        except (OSError, subprocess.CalledProcessError):
            return None

        persisted_path = Path(tempfile.gettempdir()) / f"medol-topology-{abs(hash(tuple(edges)))}.png"
        persisted_path.write_bytes(image_path.read_bytes())
        return persisted_path


def escape_dot_label(value):
    return value.replace("\\", "\\\\").replace('"', '\\"')


def add_run(paragraph, text):
    run = paragraph.add_run(text)
    run.font.name = BODY_FONT
    run.font.size = Pt(10.5)
    set_run_font(run, BODY_FONT)


def clean_inline(value):
    return (
        value.replace("**", "")
        .replace("\\|", "|")
        .replace("&lt;", "<")
        .replace("&gt;", ">")
        .strip()
    )


def set_style_font(style, font_name):
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    fonts.set(qn("w:ascii"), font_name)
    fonts.set(qn("w:hAnsi"), font_name)
    fonts.set(qn("w:eastAsia"), font_name)
    fonts.set(qn("w:cs"), font_name)


def set_run_font(run, font_name):
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    fonts.set(qn("w:ascii"), font_name)
    fonts.set(qn("w:hAnsi"), font_name)
    fonts.set(qn("w:eastAsia"), font_name)
    fonts.set(qn("w:cs"), font_name)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


if __name__ == "__main__":
    main()
