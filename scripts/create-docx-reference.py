#!/usr/bin/env python3
"""Create the sample DOCX reference template used by MEDOL Word export."""

from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


def set_run_fonts(run, east_asia: str, ascii_font: str | None = None) -> None:
    ascii_font = ascii_font or east_asia
    run.font.name = ascii_font
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), ascii_font)
    rfonts.set(qn("w:hAnsi"), ascii_font)
    rfonts.set(qn("w:eastAsia"), east_asia)
    rfonts.set(qn("w:cs"), ascii_font)


def set_style_font(style, east_asia: str, ascii_font: str, size_pt: float, color: str | None = None, bold: bool | None = None) -> None:
    font = style.font
    font.name = ascii_font
    font.size = Pt(size_pt)
    if color:
        font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        font.bold = bold
    rpr = style._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), ascii_font)
    rfonts.set(qn("w:hAnsi"), ascii_font)
    rfonts.set(qn("w:eastAsia"), east_asia)
    rfonts.set(qn("w:cs"), ascii_font)


def set_paragraph_spacing(style, before: int = 0, after: int = 120, line: int = 360) -> None:
    ppr = style._element.get_or_add_pPr()
    spacing = ppr.find(qn("w:spacing"))
    if spacing is None:
        spacing = OxmlElement("w:spacing")
        ppr.append(spacing)
    spacing.set(qn("w:before"), str(before))
    spacing.set(qn("w:after"), str(after))
    spacing.set(qn("w:line"), str(line))
    spacing.set(qn("w:lineRule"), "auto")


def configure_styles(doc: Document) -> None:
    styles = doc.styles

    normal = styles["Normal"]
    set_style_font(normal, "Songti SC", "Times New Roman", 10.5, "000000")
    set_paragraph_spacing(normal, before=0, after=120, line=360)
    normal.paragraph_format.first_line_indent = Pt(21)

    body_text = styles["Body Text"]
    set_style_font(body_text, "Songti SC", "Times New Roman", 10.5, "000000")
    set_paragraph_spacing(body_text, before=0, after=120, line=360)

    title = styles["Title"]
    set_style_font(title, "Heiti SC", "Arial", 24, "000000", True)
    set_paragraph_spacing(title, before=0, after=160, line=320)
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

    subtitle = styles["Subtitle"]
    set_style_font(subtitle, "PingFang SC", "Arial", 14, "666666", False)
    set_paragraph_spacing(subtitle, before=0, after=240, line=300)
    subtitle.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

    h1 = styles["Heading 1"]
    set_style_font(h1, "Heiti SC", "Arial", 16, "000000", True)
    set_paragraph_spacing(h1, before=240, after=160, line=320)

    h2 = styles["Heading 2"]
    set_style_font(h2, "Heiti SC", "Arial", 14, "000000", True)
    set_paragraph_spacing(h2, before=200, after=120, line=320)

    h3 = styles["Heading 3"]
    set_style_font(h3, "Heiti SC", "Arial", 12, "333333", True)
    set_paragraph_spacing(h3, before=160, after=80, line=300)

    for style_name in ["TOC Heading", "Caption", "Quote", "Intense Quote"]:
        if style_name in styles:
            set_style_font(styles[style_name], "Songti SC", "Times New Roman", 10.5, "333333")
            set_paragraph_spacing(styles[style_name], before=80, after=80, line=320)

    if "Source Code" in styles:
        set_style_font(styles["Source Code"], "Menlo", "Consolas", 9, "222222")


def set_table_style(doc: Document) -> None:
    style = doc.styles["Table Grid"]
    tbl_pr = style._element.find(qn("w:tblPr"))
    if tbl_pr is None:
        tbl_pr = OxmlElement("w:tblPr")
        style._element.append(tbl_pr)
    tbl_cell_mar = tbl_pr.find(qn("w:tblCellMar"))
    if tbl_cell_mar is None:
        tbl_cell_mar = OxmlElement("w:tblCellMar")
        tbl_pr.append(tbl_cell_mar)
    for name, width in [("top", "80"), ("bottom", "80"), ("start", "120"), ("end", "120")]:
        node = tbl_cell_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tbl_cell_mar.append(node)
        node.set(qn("w:w"), width)
        node.set(qn("w:type"), "dxa")


def add_field(paragraph, instruction: str) -> None:
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    run = paragraph.add_run()
    run._r.append(begin)

    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" {instruction} "
    run = paragraph.add_run()
    run._r.append(instr)

    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    run = paragraph.add_run()
    run._r.append(separate)

    text = paragraph.add_run("1")
    set_run_fonts(text, "Songti SC", "Times New Roman")

    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run = paragraph.add_run()
    run._r.append(end)


def configure_page(doc: Document) -> None:
    section = doc.sections[0]
    section.start_type = WD_SECTION.NEW_PAGE
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.8)
    section.right_margin = Cm(2.8)
    section.header_distance = Cm(1.25)
    section.footer_distance = Cm(1.25)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    header.paragraph_format.space_after = Pt(0)
    run = header.add_run("MEDOL 文档导出")
    set_run_fonts(run, "PingFang SC", "Arial")
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(102, 102, 102)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.space_before = Pt(0)
    footer.paragraph_format.space_after = Pt(0)
    prefix = footer.add_run("第 ")
    set_run_fonts(prefix, "Songti SC", "Times New Roman")
    prefix.font.size = Pt(9)
    add_field(footer, "PAGE")
    mid = footer.add_run(" 页 / 共 ")
    set_run_fonts(mid, "Songti SC", "Times New Roman")
    mid.font.size = Pt(9)
    add_field(footer, "SECTIONPAGES")
    suffix = footer.add_run(" 页")
    set_run_fonts(suffix, "Songti SC", "Times New Roman")
    suffix.font.size = Pt(9)


def build_reference_docx(output_path: Path) -> None:
    doc = Document()
    configure_page(doc)
    configure_styles(doc)
    set_table_style(doc)

    title = doc.add_paragraph("MEDOL 中文正式文档 Reference DOCX")
    title.style = doc.styles["Title"]
    subtitle = doc.add_paragraph("本文件内容会被 Pandoc 忽略；样式、页眉页脚和页面设置会作为导出基准。")
    subtitle.style = doc.styles["Subtitle"]
    doc.add_paragraph("正文样式示例：宋体，五号，1.5 倍行距，段后 6 磅。")
    doc.add_heading("一级标题样式", level=1)
    doc.add_heading("二级标题样式", level=2)
    doc.add_heading("三级标题样式", level=3)
    table = doc.add_table(rows=2, cols=3)
    table.style = "Table Grid"
    table.rows[0].cells[0].text = "字段"
    table.rows[0].cells[1].text = "类型"
    table.rows[0].cells[2].text = "说明"
    table.rows[1].cells[0].text = "example"
    table.rows[1].cells[1].text = "String"
    table.rows[1].cells[2].text = "表格网格样式示例"

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)


def main() -> None:
    output_path = Path(sys.argv[1]) if len(sys.argv) > 1 else Path("templates/docx/zh-formal.reference.docx")
    build_reference_docx(output_path)
    print(output_path)


if __name__ == "__main__":
    main()
